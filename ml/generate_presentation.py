"""
CardioTwin AI - Project Presentation Generator (PPTX & PDF)
Digital Twin Challenge 2026 - Happiest Health
Author: Madabattula Tejo Hasini (Aditya University)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import letter, landscape
from reportlab.pdfgen import canvas
from reportlab.lib import colors

SLIDES_DATA = [
    {
        "title": "CardioTwin AI: Virtual Cardiovascular Patient Laboratory",
        "subtitle": "Digital Twin Challenge 2026 organized by Happiest Health",
        "category": "Title Slide",
        "bullets": [
            "Author: Madabattula Tejo Hasini (Solo Student Entry)",
            "Institution: Aditya University, Andhra Pradesh, India",
            "Contact: 24b11ai223@adityauniversity.in",
            "Core Focus: In-silico cardiovascular digital twins fusing EHR & wearable telemetry",
            "License: MIT Open-Source License • Local Zero-Cloud Execution"
        ]
    },
    {
        "title": "Problem Statement: The Post-Discharge Monitoring Blindspot",
        "subtitle": "Fragmented episodic clinic visits vs continuous home deterioration",
        "category": "Clinical Background",
        "bullets": [
            "Cardiovascular diseases (CVDs) remain the leading cause of global mortality.",
            "Traditional healthcare relies on episodic 15-minute consultations every 3 to 6 months.",
            "Acute autonomic destabilization often transpires within hours at home without warning.",
            "Ordinary consumer wearables record data, but lack clinical context from baseline EHR records.",
            "Need: A personalized computational avatar that fuses static clinical history with dynamic physiological sensors."
        ]
    },
    {
        "title": "Why a Cardiovascular Digital Twin?",
        "subtitle": "Transforming passive telemetry into active predictive foresight",
        "category": "Concept & Innovation",
        "bullets": [
            "Personalized Computational Avatar: Mirrors individual patient physiology in software.",
            "Continuous Multi-Modal Fusion: Combines static EHR risk factors with high-frequency wearable streams.",
            "Dynamic 24-Hour Horizon: Continuously updates deterioration risk when new readings arrive.",
            "Counterfactual 'What-If' Simulation: Allows clinicians to simulate interventions (e.g. sleep improvement, BP control) without patient harm.",
            "Auditable Decision Trail: Every inference and sensor reading is permanently logged."
        ]
    },
    {
        "title": "Target Healthcare Use Case & Prediction Scope",
        "subtitle": "Synthesized 24-hour deterioration proxy under strictly causal lookback",
        "category": "Clinical Specification",
        "bullets": [
            "Target Definition: Composite physiological deterioration over the subsequent 24-hour window.",
            "Trigger Criteria: Sustained tachycardia (>=95 bpm), parasympathetic collapse (HRV RMSSD <=18 ms), or acute hypertensive surge (SBP >=165 mmHg).",
            "Strict Zero-Leakage Policy: No future observations utilized in feature extraction at prediction time T.",
            "Scientific Honesty Disclaimer: Prototype evaluates a synthetic proxy; does not predict heart attacks or replace clinical judgment.",
            "Outcome Prevalence: 29.0% positive window occurrence across the 5,550 observation cohort."
        ]
    },
    {
        "title": "End-to-End System Architecture",
        "subtitle": "Modular, maintainable, and completely reproducible local pipeline",
        "category": "Technical Architecture",
        "bullets": [
            "Data Pipeline: Synthetic EHR Generator + 14-day multi-frequency wearable telemetry simulator.",
            "Feature Engineering: Causal rolling temporal windows (6h mean HR, 24h min HRV, sleep debt, rate-pressure index).",
            "Machine Learning Engine: Tuned Random Forest (150 trees) vs Standardized Logistic Regression baseline.",
            "Persistence Layer: Embedded SQLite database logging patients, telemetry, predictions, and simulation events.",
            "Backend: High-performance Python FastAPI service with Pydantic schema validation.",
            "Frontend: Responsive React + TypeScript dashboard with Recharts and Tailwind CSS."
        ]
    },
    {
        "title": "Static Baseline EHR Ingestion",
        "subtitle": "Anatomical and chronic vascular baseline parameters",
        "category": "Data Engineering",
        "bullets": [
            "Demographics: Age (38-80 yrs), Biological Sex, Body Mass Index (BMI).",
            "Comorbidities: Clinically correlated Hypertension, Type 2 Diabetes, Smoking Status, Prior MI history.",
            "Hemodynamic & Lipid Markers: Baseline Systolic/Diastolic BP (mmHg), Total Cholesterol, HDL, LDL (mg/dL).",
            "Renal Filtration Markers: Serum Creatinine (mg/dL) and CKD-EPI estimated eGFR (mL/min).",
            "Provenance & Integrity: Seed 42 determinism, missingness indicators, and audit creation timestamps."
        ]
    },
    {
        "title": "Dynamic Wearable Telemetry Streams",
        "subtitle": "Time-indexed physiological signals captured at 3-hour intervals",
        "category": "Data Engineering",
        "bullets": [
            "Continuous Heart Rate: Instantaneous pulse rate in beats per minute (bpm).",
            "Heart Rate Variability (HRV): Root Mean Square of Successive Differences (RMSSD in ms) reflecting parasympathetic autonomic tone.",
            "Sleep Architecture: Daily nocturnal sleep duration (hours) and actigraphy sleep quality index (0-100).",
            "Physical Activity: Interval step count and ambulatory movement intensity.",
            "Ambulatory Blood Pressure: Physiologically correlated estimated systolic BP.",
            "Simulated Loss: Realistic 2.0% packet dropout rate with causal forward-fill imputation."
        ]
    },
    {
        "title": "Machine Learning Pipeline & Zero-Leakage Validation",
        "subtitle": "Patient-level grouped holdout evaluation with zero cross-patient overlap",
        "category": "Machine Learning",
        "bullets": [
            "Partitioning: GroupShuffleSplit on unique patient IDs (40 training patients, 10 held-out test patients).",
            "Zero Data Leakage: Verified 0 overlapping patients between train and test splits.",
            "StandardScaler Isolation: Preprocessing fitted exclusively on the 40 training patients.",
            "Baseline Model: Logistic Regression with L2 regularization and balanced class weights.",
            "Comparison Model: Random Forest Classifier (150 estimators, max_depth=6, min_samples_leaf=4).",
            "Explainability: Gini impurity ranking and local feature deviation vectors."
        ]
    },
    {
        "title": "Empirical Evaluation Results (Held-Out Test Cohort)",
        "subtitle": "Genuine executed metrics from saved artifacts - zero fabricated numbers",
        "category": "Results & Validation",
        "bullets": [
            "ROC-AUC: Random Forest achieved 0.9367 vs Logistic Regression 0.8217 (+11.5% gain).",
            "Average Precision (PR-AUC): Random Forest achieved 0.8939 vs Logistic Regression 0.8233.",
            "Balanced F1-Score: Random Forest achieved 0.8295 vs Logistic Regression 0.6733.",
            "Recall (Sensitivity): Random Forest detected 93.29% of deterioration windows (vs 85.53% LR).",
            "Brier Calibration Loss: 0.1221 (RF) vs 0.2373 (LR) demonstrating superior probability reliability.",
            "Confusion Matrix (N=1,110 test samples): TN=482, FP=151, FN=32, TP=445."
        ]
    },
    {
        "title": "Interactive Prototype: Research Dashboard & What-If Studio",
        "subtitle": "Six connected screens built for clinical researchers",
        "category": "User Experience",
        "bullets": [
            "Screen A (Overview): Cohort-wide risk distribution, active telemetry streams, and patient explorer.",
            "Screen B (Patient Twin): Static EHR baseline cards vs dynamic Recharts (HR, HRV RMSSD ms, Sleep, Steps).",
            "Screen C (Simulation & What-If): Timeline advancer, stress spike injection, and counterfactual parameter sliders.",
            "Screen D (Evaluation): Genuine ROC/PR metrics, confusion matrices, and calibration disclosure.",
            "Screen E (Explainability): Gini feature rankings, data quality monitors, and correlation vs causation alerts.",
            "Screen F (About): Competition background, student details, and ethical safeguards."
        ]
    },
    {
        "title": "Safety, Ethical Safeguards & Medical Disclaimers",
        "subtitle": "Scientific honesty and strict adherence to competition boundaries",
        "category": "Ethics & Safety",
        "bullets": [
            "100% Synthetic Data: Zero real patient records or identifiable health information (PHI) utilized.",
            "No Medical Overclaiming: Explicitly disclaims diagnosis, treatment prescription, or heart attack prediction.",
            "Correlation != Causation: High feature importance denotes algorithmic association, not biological proof.",
            "Prominent Dashboard Banners: Unmissable research prototype disclaimers on every screen.",
            "Local Confidentiality: No telemetry or patient IDs are ever uploaded to cloud or commercial APIs."
        ]
    },
    {
        "title": "Expected Impact, Future Roadmap & Conclusion",
        "subtitle": "Bridging academic digital-twin prototypes with future clinical validation",
        "category": "Conclusion & Next Steps",
        "bullets": [
            "Clinical Trial Roadmap: Prospective multi-center validation with institutional review board (IRB) approval.",
            "Sensor Benchmark: Cross-validating wearable PPG sensors against gold-standard 12-lead Holter ECGs.",
            "Demographic Tuning: Calibrating distributions against representative Indian population cardiovascular registries.",
            "Regulatory Compliance: Building SaMD audit documentation compliant with CDSCO and FDA guidelines.",
            "Summary: CardioTwin AI delivers a fully working, reproducible, technically credible digital-twin prototype."
        ]
    }
]

def generate_pptx():
    os.makedirs("presentation", exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    for slide_info in SLIDES_DATA:
        slide = prs.slides.add_slide(blank_layout)
        
        # Background shape (Dark Navy)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(7, 13, 27)
        bg.line.color.rgb = RGBColor(7, 13, 27)

        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.5))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p_cat = tf.paragraphs[0]
        p_cat.text = slide_info["category"].upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = RGBColor(45, 212, 191)  # Teal
        
        p_title = tf.add_paragraph()
        p_title.text = slide_info["title"]
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(248, 250, 252)
        
        p_sub = tf.add_paragraph()
        p_sub.text = slide_info["subtitle"]
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = RGBColor(148, 163, 184)

        # Content Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.733), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 23, 42)
        card.line.color.rgb = RGBColor(51, 65, 85)
        card.line.width = Pt(1.5)

        # Bullets inside Card
        content_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.4), Inches(11.1), Inches(4.0))
        c_tf = content_box.text_frame
        c_tf.word_wrap = True

        for i, bullet in enumerate(slide_info["bullets"]):
            p = c_tf.paragraphs[0] if i == 0 else c_tf.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(226, 232, 240)
            p.space_after = Pt(12)

        # Footer
        footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(11.7), Inches(0.5))
        f_tf = footer.text_frame
        f_p = f_tf.paragraphs[0]
        f_p.text = "CardioTwin AI • Digital Twin Challenge 2026 (Happiest Health) • Aditya University • Madabattula Tejo Hasini"
        f_p.font.size = Pt(9.5)
        f_p.font.color.rgb = RGBColor(100, 116, 139)

    pptx_path = "presentation/CardioTwin_AI_Presentation.pptx"
    prs.save(pptx_path)
    print(f"[OK] Saved {pptx_path}")


def generate_pdf():
    os.makedirs("presentation", exist_ok=True)
    pdf_path = "presentation/CardioTwin_AI_Presentation.pdf"
    c = canvas.Canvas(pdf_path, pagesize=landscape(letter))
    width, height = landscape(letter)

    for slide in SLIDES_DATA:
        # Background
        c.setFillColor(colors.HexColor("#070D1B"))
        c.rect(0, 0, width, height, fill=1, stroke=0)

        # Category
        c.setFillColor(colors.HexColor("#2DD4BF"))
        c.setFont("Helvetica-Bold", 10)
        c.drawString(45, height - 40, slide["category"].upper())

        # Title
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 18)
        c.drawString(45, height - 65, slide["title"])

        # Subtitle
        c.setFillColor(colors.HexColor("#94A3B8"))
        c.setFont("Helvetica", 11)
        c.drawString(45, height - 85, slide["subtitle"])

        # Card container
        c.setFillColor(colors.HexColor("#0F172A"))
        c.setStrokeColor(colors.HexColor("#334155"))
        c.roundRect(40, 50, width - 80, height - 150, 8, fill=1, stroke=1)

        # Bullets
        c.setFillColor(colors.HexColor("#E2E8F0"))
        c.setFont("Helvetica", 11)
        y_pos = height - 180
        for bullet in slide["bullets"]:
            c.drawString(65, y_pos, f"• {bullet}")
            y_pos -= 28

        # Footer
        c.setFillColor(colors.HexColor("#64748B"))
        c.setFont("Helvetica", 8)
        c.drawCentredString(width / 2.0, 25, "CardioTwin AI • Digital Twin Challenge 2026 (Happiest Health) • Aditya University • Madabattula Tejo Hasini")

        c.showPage()

    c.save()
    print(f"[OK] Saved {pdf_path}")


if __name__ == "__main__":
    generate_pptx()
    generate_pdf()
